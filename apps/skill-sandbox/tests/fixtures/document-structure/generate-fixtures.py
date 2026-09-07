"""Create the committed real OOXML/PDF parser fixtures."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).parent


def pdf() -> None:
    path = ROOT / "cross-page-table.pdf"
    canvas = Canvas(str(path), pagesize=(420, 300), pageCompression=0, invariant=1)
    for page_number, rows in [(1, [("甲", "120"), ("乙", "230")]), (2, [("丙", "340"), ("丁", "450")])]:
        x_positions = [40, 210, 380]
        y_positions = [250, 210, 170, 130]
        for x in x_positions:
            canvas.line(x, y_positions[-1], x, y_positions[0])
        for y in y_positions:
            canvas.line(x_positions[0], y, x_positions[-1], y)
        for column, value in enumerate(("Region", "Revenue")):
            canvas.drawString(x_positions[column] + 8, 225, value)
        for row_index, (region, revenue) in enumerate(rows, start=1):
            canvas.drawString(48, 225 - row_index * 40, region.encode("ascii", "replace").decode())
            canvas.drawString(218, 225 - row_index * 40, revenue)
        canvas.drawString(40, 275, f"Quarterly table - page {page_number}")
        canvas.showPage()
    canvas.save()


def docx() -> None:
    document = Document()
    document.add_paragraph("董事会摘要 Board summary")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "收入"
    table.cell(1, 1).text = "120"
    document.save(ROOT / "native-locators.docx")


def pptx() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "第一张：经营摘要"
    slide2 = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2))
    for row, values in enumerate((("指标", "数值"), ("收入", "120"))):
        for column, value in enumerate(values):
            shape.table.cell(row, column).text = value
    presentation.save(ROOT / "native-locators.pptx")


def xlsx() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet["A1"] = "指标"
    sheet["B1"] = "数值"
    sheet["A2"] = "收入"
    sheet["B2"] = 120
    sheet["C2"] = "=B2*2"
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "合并说明"
    workbook.save(ROOT / "native-locators.xlsx")


if __name__ == "__main__":
    ROOT.mkdir(parents=True, exist_ok=True)
    pdf()
    docx()
    pptx()
    xlsx()
