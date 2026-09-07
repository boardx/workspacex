#!/usr/bin/env python3
"""Extract defensible native locations from PDF and OOXML documents.

This adapter deliberately delegates parsing to maintained format libraries.  The
only cross-page inference is a conservative PDF table grouping rule: adjacent
fragments must have the same non-empty repeated header and column count.  The
result records that rule as heuristic; it never invents Word page numbers.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

LIMITS = json.loads(Path(__file__).with_name('document-structure-limits.json').read_text())
MAX_CHUNKS = LIMITS['maxChunks']


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def add(chunks: list[dict[str, Any]], chunk: dict[str, Any]) -> None:
    if len(chunks) >= MAX_CHUNKS:
        raise RuntimeError("document_structure_chunk_limit_exceeded")
    if chunk["text"]:
        chunks.append(chunk)


def parse_pdf(source: Path) -> dict[str, Any]:
    import pdfplumber

    chunks: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    previous: dict[str, Any] | None = None
    with pdfplumber.open(source) as document:
        for page_number, page in enumerate(document.pages, start=1):
            for word in page.extract_words():
                add(
                    chunks,
                    {
                        "type": "pdf_text",
                        "text": text(word.get("text")),
                        "locator": {
                            "pageNumber": page_number,
                            "bbox": [round(float(word[key]), 3) for key in ("x0", "top", "x1", "bottom")],
                        },
                    },
                )
            for page_table_index, table in enumerate(page.find_tables(), start=0):
                rows = table.extract() or []
                header = tuple(text(cell) for cell in (rows[0] if rows else []))
                column_count = max((len(row) for row in rows), default=0)
                repeated = bool(
                    previous
                    and previous["pageNumber"] == page_number - 1
                    and previous["header"] == header
                    and previous["columnCount"] == column_count
                    and any(header)
                )
                table_id = previous["tableId"] if repeated else f"pdf-table-{len(tables)}"
                fragment_index = previous["fragmentIndex"] + 1 if repeated else 0
                bbox = [round(float(value), 3) for value in table.bbox]
                tables.append(
                    {
                        "tableId": table_id,
                        "fragmentIndex": fragment_index,
                        "pageNumber": page_number,
                        "pageTableIndex": page_table_index,
                        "bbox": bbox,
                        "columnCount": column_count,
                        "continuationDetection": "repeated_header_and_columns" if repeated else "none",
                    }
                )
                for row_index, row in enumerate(table.rows):
                    for column_index, cell_bbox in enumerate(row.cells):
                        if cell_bbox is None:
                            continue
                        cell_text = text(page.crop(cell_bbox).extract_text())
                        add(
                            chunks,
                            {
                                "type": "pdf_table_cell",
                                "text": cell_text,
                                "locator": {
                                    "tableId": table_id,
                                    "fragmentIndex": fragment_index,
                                    "pageNumber": page_number,
                                    "pageTableIndex": page_table_index,
                                    "rowIndex": row_index,
                                    "columnIndex": column_index,
                                    "bbox": [round(float(value), 3) for value in cell_bbox],
                                },
                            },
                        )
                previous = {
                    "tableId": table_id,
                    "fragmentIndex": fragment_index,
                    "pageNumber": page_number,
                    "header": header,
                    "columnCount": column_count,
                }
    if not chunks:
        raise RuntimeError("document_structure_empty")
    return {
        "engine": {"name": "pdfplumber", "version": pdfplumber.__version__},
        "sourceFormat": "pdf",
        "coordinateSpace": "pdf_points_top_left",
        "tables": tables,
        "chunks": chunks,
    }


def parse_docx(source: Path) -> dict[str, Any]:
    import docx

    document = docx.Document(source)
    chunks: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        add(
            chunks,
            {
                "type": "docx_paragraph",
                "text": text(paragraph.text),
                "locator": {"paragraphIndex": paragraph_index},
            },
        )
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                add(
                    chunks,
                    {
                        "type": "docx_table_cell",
                        "text": text(cell.text),
                        "locator": {
                            "tableIndex": table_index,
                            "rowIndex": row_index,
                            "columnIndex": column_index,
                        },
                    },
                )
    if not chunks:
        raise RuntimeError("document_structure_empty")
    return {
        "engine": {"name": "python-docx", "version": docx.__version__},
        "sourceFormat": "docx",
        "coordinateSpace": "ooxml_native",
        "chunks": chunks,
    }


def parse_pptx(source: Path) -> dict[str, Any]:
    import pptx

    presentation = pptx.Presentation(source)
    chunks: list[dict[str, Any]] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for element_index, shape in enumerate(slide.shapes):
            base = {
                "slideNumber": slide_number,
                "elementIndex": element_index,
                "elementName": shape.name,
            }
            if getattr(shape, "has_text_frame", False):
                add(chunks, {"type": "pptx_element", "text": text(shape.text), "locator": base})
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows):
                    for column_index, cell in enumerate(row.cells):
                        add(
                            chunks,
                            {
                                "type": "pptx_table_cell",
                                "text": text(cell.text),
                                "locator": {
                                    **base,
                                    "rowIndex": row_index,
                                    "columnIndex": column_index,
                                },
                            },
                        )
    if not chunks:
        raise RuntimeError("document_structure_empty")
    return {
        "engine": {"name": "python-pptx", "version": pptx.__version__},
        "sourceFormat": "pptx",
        "coordinateSpace": "ooxml_native",
        "chunks": chunks,
    }


def parse_xlsx(source: Path) -> dict[str, Any]:
    import openpyxl

    workbook = openpyxl.load_workbook(source, read_only=False, data_only=False, keep_links=False)
    chunks: list[dict[str, Any]] = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            merged_by_cell: dict[str, str] = {}
            for merged in sheet.merged_cells.ranges:
                for row in sheet[merged.coord]:
                    for cell in row:
                        merged_by_cell[cell.coordinate] = merged.coord
            for row in sheet.iter_rows():
                for cell in row:
                    value = text(cell.value)
                    if not value:
                        continue
                    locator: dict[str, Any] = {
                        "sheetIndex": sheet_index,
                        "sheetName": sheet.title,
                        "address": cell.coordinate,
                        "row": cell.row,
                        "column": cell.column,
                    }
                    if cell.coordinate in merged_by_cell:
                        locator["mergedRange"] = merged_by_cell[cell.coordinate]
                    add(chunks, {"type": "xlsx_cell", "text": value, "locator": locator})
    finally:
        workbook.close()
    if not chunks:
        raise RuntimeError("document_structure_empty")
    return {
        "engine": {"name": "openpyxl", "version": openpyxl.__version__},
        "sourceFormat": "xlsx",
        "coordinateSpace": "ooxml_native",
        "chunks": chunks,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--media-type", required=True)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    parsers = {
        "application/pdf": parse_pdf,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": parse_docx,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": parse_pptx,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": parse_xlsx,
    }
    selected = parsers.get(args.media_type)
    if selected is None:
        raise RuntimeError("document_structure_format_unsupported")
    result = {"schemaVersion": 1, **selected(args.source)}
    encoded = json.dumps(result, ensure_ascii=False, separators=(",", ":")).encode('utf-8')
    if len(encoded) > LIMITS['maxOutputBytes']:
        raise RuntimeError('document_structure_output_limit_exceeded')
    args.output.write_bytes(encoded)


if __name__ == "__main__":
    main()
